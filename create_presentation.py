from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# Theme colors
PRIMARY = RGBColor(29, 78, 216)
DARK = RGBColor(15, 23, 42)
ACCENT = RGBColor(100, 116, 139)
LIGHT = RGBColor(248, 250, 252)
GREEN = RGBColor(5, 150, 105)

# Slide 1 - Title
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.background.fill.solid()
slide.background.fill.fore_color.rgb = LIGHT

# Add title text
title = slide.shapes.title
title.text = 'FinSight AI'
title.text_frame.paragraphs[0].font.size = Pt(28)
title.text_frame.paragraphs[0].font.bold = True
title.text_frame.paragraphs[0].font.color.rgb = PRIMARY

subtitle = slide.placeholders[1]
subtitle.text = 'Intelligent Stock Market Prediction & Investment Assistant'
subtitle.text_frame.paragraphs[0].font.size = Pt(18)
subtitle.text_frame.paragraphs[0].font.color.rgb = DARK

# Add accent bar
shape = slide.shapes.add_shape(1, Inches(0.7), Inches(2.4), Inches(11.9), Inches(0.1))
shape.fill.solid()
shape.fill.fore_color.rgb = PRIMARY
shape.line.color.rgb = PRIMARY

# Slide 2 - Overview
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Project Overview'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    'FinSight AI is a stock prediction and analysis web application built in Python using Streamlit.',
    'The system helps users analyze stocks, view technical indicators, and get AI-based BUY/HOLD/SELL signals.',
    'It is designed for financial analysis, investment support, and educational use.',
    'The project follows a practical ML workflow: data collection, feature engineering, training, evaluation, and deployment.'
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.font.size = Pt(18)
    p.font.color.rgb = DARK

# Slide 3 - Problem & Objective
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Problem Statement & Objective'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    'Problem: Stock markets are volatile and difficult to interpret manually using raw price movement alone.',
    'Objective: Build an AI-powered assistant that predicts next-day price movement and supports decision-making.',
    'Goal: Provide a user-friendly dashboard with technical indicators, market insights, and model explainability.',
    'Impact: Helps beginners and analysts understand stock trends using data-driven signals.'
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.font.size = Pt(18)
    p.font.color.rgb = DARK

# Slide 4 - Methodology
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Methodology'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    '1. Collect historical stock data from Yahoo Finance.',
    '2. Perform feature engineering using OHLCV data and technical indicators.',
    '3. Train a Random Forest Regressor for next-day close price prediction.',
    '4. Evaluate the model using MAE, RMSE, and R² score.',
    '5. Deploy the model through an interactive Streamlit dashboard.'
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.font.size = Pt(18)
    p.font.color.rgb = DARK

# Slide 5 - Technical Features
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Key Technical Features'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    'Moving Averages (MA10, MA20)',
    'Relative Strength Index (RSI)',
    'MACD and MACD Histogram',
    'Historical Volatility',
    'Daily Return and Price Momentum Signals'
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.font.size = Pt(18)
    p.font.color.rgb = DARK

# Slide 6 - Model Architecture
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Model & Prediction Workflow'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    'Input Data → Feature Engineering → Normalization → Random Forest Training',
    'Predicted Close Price → Signal Generation → BUY / HOLD / SELL Decision',
    'Model stores: trained estimator, scaler, and metadata files for reuse.',
    'Dashboard displays live market status, company metrics, and ML-based insights.'
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.font.size = Pt(18)
    p.font.color.rgb = DARK

# Slide 7 - Model Performance
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Model Performance Snapshot'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    'Random Forest Regressor used for prediction.',
    'Training Set: MAE = ₹15.65, RMSE = ₹19.96, R² = 0.9781',
    'Test Set: MAE = ₹34.10, RMSE = ₹46.73, R² = 0.8476',
    'These metrics indicate strong trend-learning ability and practical usefulness for stock prediction.'
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.font.size = Pt(18)
    p.font.color.rgb = DARK

# Slide 8 - Conclusion
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = 'Conclusion'
body = slide.shapes.placeholders[1].text_frame
body.clear()
for i, text in enumerate([
    'FinSight AI combines machine learning, technical analysis, and interactive UI design.',
    'It delivers a practical stock forecasting assistant that is easy for users to understand and use.',
    'The system can be extended with more companies, additional indicators, and improved forecasting models.',
    'Future scope: deeper AI insights, real-time market ingestion, and portfolio analysis.'
]):
    p = body.paragraphs[0] if i == 0 else body.add_paragraph()
    p.text = text
    p.level = 0
    p.font.size = Pt(18)
    p.font.color.rgb = DARK

# Save presentation
prs.save('FinSight_AI_Project_Presentation.pptx')
print('Presentation created: FinSight_AI_Project_Presentation.pptx')
